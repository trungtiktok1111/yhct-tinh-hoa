# -*- coding: utf-8 -*-
"""
Trích xuất ảnh minh họa vị thuốc từ slide "intro" (slide có "Tên:" / "Họ:")
trong các file PPTX của DƯỢC CỔ TRUYỀN. Lưu ảnh ra public/images/herbs/
và ghi map {vn_name, file, slide} -> local path để script TS đọc và cập nhật DB.
"""
import glob, json, os, re, sys, hashlib
from pptx import Presentation

ROOT = r"D:\Tài liệu đông y\tài liệu y học\DƯỢC CỔ TRUYỀN"
OUT_DIR = r"D:\yhct-tinh-hoa\public\images\herbs"
os.makedirs(OUT_DIR, exist_ok=True)

def slug(name, salt):
    base = re.sub(r"[^a-zA-Z0-9]+", "-", name.strip().lower()).strip("-")
    h = hashlib.md5(salt.encode("utf-8")).hexdigest()[:6]
    return f"{base}-{h}"

def slide_texts(slide):
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            out.append(shape.text_frame.text.strip())
    return out

def is_intro_slide(texts):
    return any(t.startswith("Tên:") for t in texts)

def parse_vn_name(texts):
    for t in texts:
        if not t.startswith("Tên:") and not t.startswith("Họ:"):
            lines = [l.strip() for l in t.split("\n") if l.strip()]
            if lines:
                return lines[0]
    return None

def main():
    files = sorted(glob.glob(os.path.join(ROOT, "0*.pptx"))) + sorted(glob.glob(os.path.join(ROOT, "1*.pptx")))
    mapping = []
    for f in files:
        fname = os.path.basename(f)
        try:
            p = Presentation(f)
        except Exception as e:
            print(f"{fname}: ERROR open {e}", file=sys.stderr)
            continue
        slides = list(p.slides)
        for i, slide in enumerate(slides):
            texts = slide_texts(slide)
            if not is_intro_slide(texts):
                continue
            vn_name = parse_vn_name(texts)
            if not vn_name:
                continue
            pic_shape = None
            for shape in slide.shapes:
                if shape.shape_type == 13:  # PICTURE
                    pic_shape = shape
                    break
            if pic_shape is None:
                continue
            try:
                image = pic_shape.image
                ext = image.ext
                fname_slug = slug(vn_name, f"{fname}-{i}")
                out_path = os.path.join(OUT_DIR, f"{fname_slug}.{ext}")
                with open(out_path, "wb") as fh:
                    fh.write(image.blob)
                mapping.append({
                    "vn_name": vn_name,
                    "file": fname,
                    "slide": i + 1,
                    "image_path": f"/images/herbs/{fname_slug}.{ext}",
                })
            except Exception as e:
                print(f"{fname} slide {i+1}: image extract error {e}", file=sys.stderr)
        print(f"{fname}: done", file=sys.stderr)

    with open(r"D:\yhct-tinh-hoa\scripts\herb_images_map.json", "w", encoding="utf-8") as fh:
        json.dump(mapping, fh, ensure_ascii=False, indent=1)
    print(f"TOTAL images extracted: {len(mapping)}", file=sys.stderr)

if __name__ == "__main__":
    main()
