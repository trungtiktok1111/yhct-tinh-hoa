# -*- coding: utf-8 -*-
"""Phân tích file 'CHAM CUU THEO CHUNG.pptx' (Xử phương châm cứu cổ điển): mỗi
chứng/triệu chứng có 1 slide bảng tổng quan (Công thức/Xuất xứ/Công năng/Chủ trị)
và 1 slide bảng chi tiết Huyệt/Ý nghĩa — ghép theo TIÊU ĐỀ slide (không theo vị
trí) để tránh lệch khi có slide dư (vd công thức dài chiếm nhiều slide). Chỉ lấy
đúng nội dung có trong slide, không bịa thêm."""
import json, sys
from pptx import Presentation

PATH = r"D:\Tài liệu đông y\tài liệu y học\CHÂM CỨU\CHAM CUU THEO CHUNG.pptx"

def slide_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip().split("\n")[0].strip()
    return None

def slide_table(slide):
    for sh in slide.shapes:
        if getattr(sh, "has_table", False):
            return [[c.text.strip() for c in row.cells] for row in sh.table.rows]
    return None

def main():
    p = Presentation(PATH)
    slides = list(p.slides)
    overviews = []  # list of (title, overview dict, slide_no)
    details = {}  # title -> list of point dict (first one found kept)
    for idx, s in enumerate(slides):
        title = slide_title(s)
        table = slide_table(s)
        if table is None or not title:
            continue
        header0 = table[0][0] if table[0] else ""
        if header0 == "Huyệt":
            header = table[0]
            rows = []
            for row in table[1:]:
                rows.append({header[j]: row[j] for j in range(min(len(row), len(header)))})
            details.setdefault(title, rows)
        else:
            overview = {row[0]: row[1] for row in table if len(row) >= 2 and row[0]}
            overviews.append({"name": title, "overview": overview, "slide": idx + 1})

    items = []
    for ov in overviews:
        items.append({
            "name": ov["name"],
            "overview": ov["overview"],
            "points_detail": details.get(ov["name"], []),
            "slide": ov["slide"],
        })
    with open(r"D:\yhct-tinh-hoa\scripts\cham_cuu_theo_chung_parsed.json", "w", encoding="utf-8") as fh:
        json.dump(items, fh, ensure_ascii=False, indent=1)
    missing = [it["name"] for it in items if not it["points_detail"]]
    print(f"TOTAL: {len(items)}  missing_detail: {missing}", file=sys.stderr)

if __name__ == "__main__":
    main()
