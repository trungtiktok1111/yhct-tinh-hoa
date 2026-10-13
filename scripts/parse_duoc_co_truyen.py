# -*- coding: utf-8 -*-
"""
Phân tích các file PPTX trong DƯỢC CỔ TRUYỀN theo mẫu 4-slide/vị thuốc:
  Slide A (intro): "Tên: <latin>", "Họ: <family>", <tên VN>, <tên latin dược liệu>
  Slide B: bảng "Đặc điểm" -> Tính, Vị, Quy kinh, Hướng tác động, Vị trí tác động
  Slide C: bảng "Tác dụng YHCT" -> (tác dụng, ghi chú)
  Slide D: bảng "Sử dụng" -> Liều lượng, Thời gian sắc, Độc tính, Chống chỉ định, Thai kỳ
Xuất ra JSON có cấu trúc, không suy diễn gì ngoài những gì đọc được trong slide.
"""
import glob, json, os, re, sys
from pptx import Presentation

ROOT = r"D:\Tài liệu đông y\tài liệu y học\DƯỢC CỔ TRUYỀN"

def slide_shapes_text(slide):
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            out.append(("text", shape.text_frame.text.strip()))
        elif shape.has_table:
            rows = []
            for row in shape.table.rows:
                rows.append([c.text.strip() for c in row.cells])
            out.append(("table", rows))
    return out

def is_intro_slide(shapes):
    for kind, val in shapes:
        if kind == "text" and val.startswith("Tên:") :
            return True
    return False

def parse_intro(shapes):
    latin_genus = None
    family = None
    vn_name = None
    latin_pharm = None
    for kind, val in shapes:
        if kind != "text":
            continue
        lines = [l.strip() for l in val.split("\n") if l.strip()]
        for l in lines:
            if l.startswith("Tên:"):
                latin_genus = l[len("Tên:"):].strip()
            elif l.startswith("Họ:"):
                family = l[len("Họ:"):].strip()
        # the block without Tên:/Họ: prefix containing 2 lines is VN name + latin pharm name
        if not val.startswith("Tên:") and not val.startswith("Họ:"):
            if len(lines) >= 2 and not vn_name:
                vn_name, latin_pharm = lines[0], lines[1]
            elif len(lines) == 1 and not vn_name:
                vn_name = lines[0]
    return vn_name, latin_pharm, latin_genus, family

def parse_dac_diem_table(rows):
    tinh = []
    vi = []
    quy_kinh = []
    huong_tac_dong = []
    vi_tri_tac_dong = []
    current_label = None
    for row in rows[1:]:
        if len(row) < 2:
            continue
        label = row[0].strip()
        value = row[1].strip()
        if label:
            current_label = label
        if current_label == "Tính" and value:
            tinh.append(value)
        elif current_label == "Vị" and value:
            vi.append(value)
        elif current_label == "Quy kinh" and value:
            quy_kinh.append(value)
        elif current_label == "Hướng tác động" and value:
            huong_tac_dong.append(value)
        elif current_label == "Vị trí tác động" and value:
            vi_tri_tac_dong.append(value)
    return tinh, vi, quy_kinh, huong_tac_dong, vi_tri_tac_dong

def parse_tac_dung_table(rows):
    actions = []
    for row in rows[1:]:
        if not row or not row[0].strip():
            continue
        name = row[0].strip()
        note = row[1].strip() if len(row) > 1 else ""
        actions.append({"action": name, "note": note})
    return actions

def parse_su_dung_table(rows):
    d = {}
    for row in rows[1:]:
        if len(row) < 2:
            continue
        label = row[0].strip()
        value = row[1].strip()
        if label:
            d[label] = value
    return d

def parse_file(path):
    p = Presentation(path)
    slides = list(p.slides)
    herbs = []
    i = 0
    n = len(slides)
    while i < n:
        shapes = slide_shapes_text(slides[i])
        if is_intro_slide(shapes):
            vn_name, latin_pharm, latin_genus, family = parse_intro(shapes)
            herb = {
                "vn_name": vn_name,
                "latin_pharm_name": latin_pharm,
                "latin_genus": latin_genus,
                "family": family,
                "tinh": [], "vi": [], "quy_kinh": [], "huong_tac_dong": [], "vi_tri_tac_dong": [],
                "tac_dung": [],
                "lieu_luong": None, "thoi_gian_sac": None, "doc_tinh": None,
                "chong_chi_dinh": None, "thai_ky": None,
                "slide_start": i + 1,
            }
            j = i + 1
            while j < n:
                s2 = slide_shapes_text(slides[j])
                if is_intro_slide(s2):
                    break
                has_table = False
                for kind, val in s2:
                    if kind == "table":
                        has_table = True
                        header = val[0][0] if val and val[0] else ""
                        if "Đặc điểm" in header or (len(val[0]) > 0 and "Đặc điểm" in " ".join(val[0])):
                            tinh, vi, qk, htd, vtd = parse_dac_diem_table(val)
                            herb["tinh"] += tinh; herb["vi"] += vi; herb["quy_kinh"] += qk
                            herb["huong_tac_dong"] += htd; herb["vi_tri_tac_dong"] += vtd
                        elif val and val[0] and "Tác dụng" in " ".join(val[0]):
                            herb["tac_dung"] += parse_tac_dung_table(val)
                        elif val and val[0] and "Sử dụng" in " ".join(val[0]):
                            sd = parse_su_dung_table(val)
                            herb["lieu_luong"] = sd.get("Liều lượng") or herb["lieu_luong"]
                            herb["thoi_gian_sac"] = sd.get("Thời gian sắc") or herb["thoi_gian_sac"]
                            herb["doc_tinh"] = sd.get("Độc tính") or herb["doc_tinh"]
                            herb["chong_chi_dinh"] = sd.get("Chống chỉ định") or herb["chong_chi_dinh"]
                            herb["thai_ky"] = sd.get("Thai kỳ") or herb["thai_ky"]
                j += 1
            herb["slide_end"] = j
            if herb["vn_name"]:
                herbs.append(herb)
            i = j
        else:
            i += 1
    return herbs

def main():
    files = sorted(glob.glob(os.path.join(ROOT, "0*.pptx"))) + sorted(glob.glob(os.path.join(ROOT, "1*.pptx")))
    all_data = {}
    for f in files:
        name = os.path.basename(f)
        try:
            herbs = parse_file(f)
            all_data[name] = herbs
            print(f"{name}: {len(herbs)} vị thuốc", file=sys.stderr)
        except Exception as e:
            print(f"{name}: ERROR {e}", file=sys.stderr)
            all_data[name] = {"error": str(e)}
    out_path = r"D:\yhct-tinh-hoa\scripts\duoc_co_truyen_parsed.json"
    with open(out_path, "w", encoding="utf-8") as fh:
        json.dump(all_data, fh, ensure_ascii=False, indent=1)
    total = sum(len(v) for v in all_data.values() if isinstance(v, list))
    print(f"TOTAL herbs parsed: {total}", file=sys.stderr)

if __name__ == "__main__":
    main()
