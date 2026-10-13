# -*- coding: utf-8 -*-
"""Phân tích file 'CAC HUYET THUONG DUNG.pptx': mỗi slide = 1 huyệt
(tên huyệt + text bắt đầu bằng 'Vị trí:'). Chỉ lấy đúng những gì có trong slide."""
import json, re, sys
from pptx import Presentation

PATH = r"D:\Tài liệu đông y\tài liệu y học\CHÂM CỨU\CAC HUYET THUONG DUNG.pptx"

def main():
    p = Presentation(PATH)
    slides = list(p.slides)
    points = []
    current_section = None
    for i, slide in enumerate(slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
        cleaned = [t for t in texts if not re.fullmatch(r"\d+", t)]
        if not cleaned:
            continue
        loc = next((t for t in cleaned if t.startswith("Vị trí")), None)
        names = [t for t in cleaned if t is not loc]
        if not names:
            continue
        name_text = names[0]
        if loc is None and len(cleaned) == 1:
            # slide chỉ có 1 dòng text, không có "Vị trí" -> có thể là tiêu đề mục (vd "Các huyệt vùng đầu mặt")
            current_section = name_text
            continue
        points.append({
            "name": name_text,
            "location": loc.replace("Vị trí:", "").replace("Vị trí :", "").strip() if loc else None,
            "section": current_section,
            "slide": i + 1,
        })
    with open(r"D:\yhct-tinh-hoa\scripts\huyet_thuong_dung_parsed.json", "w", encoding="utf-8") as fh:
        json.dump(points, fh, ensure_ascii=False, indent=1)
    print(f"TOTAL huyệt: {len(points)}", file=sys.stderr)

if __name__ == "__main__":
    main()
