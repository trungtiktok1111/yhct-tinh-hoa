# -*- coding: utf-8 -*-
"""Phân loại toàn bộ PDF (text), PPTX, DOCX trong thư mục nguồn để lập kế hoạch nhập liệu."""
import glob, os, json, sys
import pypdf
from pptx import Presentation
import docx

ROOT = r"D:\Tài liệu đông y\tài liệu y học"
OUT = r"D:\yhct-tinh-hoa\scripts\master_manifest.json"

def classify_pdf(path):
    try:
        r = pypdf.PdfReader(path)
        n = len(r.pages)
        sample_n = min(n, 3)
        total = 0
        for p in range(sample_n):
            try:
                total += len(r.pages[p].extract_text() or "")
            except Exception:
                pass
        cpp = total / max(sample_n, 1)
        return {"ext": "pdf", "units": n, "unit_kind": "pages", "cpp": round(cpp, 1), "kind": "text" if cpp > 200 else "scan_or_empty"}
    except Exception as e:
        return {"ext": "pdf", "error": str(e)[:150], "kind": "error"}

def classify_pptx(path):
    try:
        p = Presentation(path)
        n = len(p.slides)
        total = 0
        for s in list(p.slides)[:5]:
            for shape in s.shapes:
                if shape.has_text_frame:
                    total += len(shape.text_frame.text)
        return {"ext": "pptx", "units": n, "unit_kind": "slides", "cpp": round(total / max(min(n,5),1), 1), "kind": "text" if total > 0 else "scan_or_empty"}
    except Exception as e:
        return {"ext": "pptx", "error": str(e)[:150], "kind": "error"}

def classify_docx(path):
    try:
        d = docx.Document(path)
        total = sum(len(p.text) for p in d.paragraphs)
        return {"ext": "docx", "units": len(d.paragraphs), "unit_kind": "paragraphs", "cpp": total, "kind": "text" if total > 50 else "scan_or_empty"}
    except Exception as e:
        return {"ext": "docx", "error": str(e)[:150], "kind": "error"}

def main():
    os.chdir(ROOT)
    pdfs = glob.glob("**/*.pdf", recursive=True)
    pptxs = glob.glob("**/*.pptx", recursive=True)
    docxs = glob.glob("**/*.docx", recursive=True) + glob.glob("**/*.doc", recursive=True)
    results = []
    all_files = [(p, "pdf") for p in pdfs] + [(p, "pptx") for p in pptxs] + [(p, "docx") for p in docxs]
    for i, (path, ext) in enumerate(all_files):
        entry = {"path": path}
        if ext == "pdf":
            entry.update(classify_pdf(path))
        elif ext == "pptx":
            entry.update(classify_pptx(path))
        else:
            entry.update(classify_docx(path))
        entry["status"] = "pending"
        results.append(entry)
        if (i + 1) % 50 == 0:
            print(f"...{i+1}/{len(all_files)}", file=sys.stderr)
    with open(OUT, "w", encoding="utf-8") as fh:
        json.dump(results, fh, ensure_ascii=False, indent=1)
    text_count = sum(1 for r in results if r.get("kind") == "text")
    print(f"TOTAL {len(results)} | text-based {text_count}")

if __name__ == "__main__":
    main()
