# -*- coding: utf-8 -*-
"""Phân tích file 'CHAM CUU BENH TANG PHU.pptx': mỗi pháp trị (vd 'Ôn bổ thận dương')
chiếm 2 slide liên tiếp — slide A: bảng Công thức/Công năng/Thao tác/Chủ trị (hoặc
Phối huyệt/Giải thích ở phần Phế), slide B: bảng chi tiết Huyệt/Ý nghĩa/Công năng.
Chỉ lấy đúng những gì có trong slide, không bịa thêm."""
import json, sys
from pptx import Presentation

PATH = r"D:\Tài liệu đông y\tài liệu y học\CHÂM CỨU\CHAM CUU BENH TANG PHU.pptx"

def slide_title(slide):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            return sh.text_frame.text.strip().split("\n")[0].strip()
    return None

def slide_table(slide):
    for sh in slide.shapes:
        if getattr(sh, "has_table", False):
            rows = [[c.text.strip() for c in row.cells] for row in sh.table.rows]
            return rows
    return None

def main():
    p = Presentation(PATH)
    slides = list(p.slides)
    section = None
    protocols = []
    i = 0
    while i < len(slides):
        title = slide_title(slides[i])
        table = slide_table(slides[i])
        if table is None:
            if title and not title.isdigit():
                section = title
            i += 1
            continue
        # slide A: tổng quan (Công thức/Nội dung hoặc Phối huyệt/Giải thích)
        overview_rows = table
        overview = {}
        for row in overview_rows:
            if len(row) >= 2 and row[0]:
                overview[row[0]] = row[1]
        # slide B kế tiếp: chi tiết huyệt
        detail_rows = []
        if i + 1 < len(slides):
            t2 = slide_table(slides[i + 1])
            if t2 and len(t2) > 0 and t2[0][0] in ("Huyệt", "Phối huyệt"):
                header = t2[0]
                for row in t2[1:]:
                    entry = {header[j]: row[j] for j in range(len(row))}
                    detail_rows.append(entry)
        protocols.append({
            "section": section,
            "name": title,
            "overview": overview,
            "points_detail": detail_rows,
            "slide_overview": i + 1,
            "slide_detail": i + 2 if detail_rows else None,
        })
        i += 2
    with open(r"D:\yhct-tinh-hoa\scripts\cham_cuu_tang_phu_parsed.json", "w", encoding="utf-8") as fh:
        json.dump(protocols, fh, ensure_ascii=False, indent=1)
    print(f"TOTAL protocols: {len(protocols)}", file=sys.stderr)

if __name__ == "__main__":
    main()
